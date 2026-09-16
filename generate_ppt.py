from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Palette ─────────────────────────────────────────────────────
BG      = RGBColor(0x0F, 0x11, 0x17)
ACCENT  = RGBColor(0x4F, 0x46, 0xE5)   # deep indigo
TEAL    = RGBColor(0x14, 0xB8, 0xA6)
AMBER   = RGBColor(0xF5, 0x9E, 0x0B)
RED     = RGBColor(0xEF, 0x44, 0x44)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT   = RGBColor(0xE5, 0xE7, 0xEB)
MUTED   = RGBColor(0x6B, 0x72, 0x80)
CARD    = RGBColor(0x1F, 0x29, 0x37)
DARK_R  = RGBColor(0x7F, 0x1D, 0x1D)   # dark red bg

W = Inches(13.33)
H = Inches(7.5)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank = prs.slide_layouts[6]

# ── Helpers ─────────────────────────────────────────────────────
def bg(slide, color=BG):
    f = slide.background.fill
    f.solid(); f.fore_color.rgb = color

def rect(slide, x, y, w, h, color, line=False):
    s = slide.shapes.add_shape(1, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = color
    if not line: s.line.fill.background()
    return s

def txt(slide, text, x, y, w, h, size=12, bold=False,
        color=WHITE, align=PP_ALIGN.LEFT, italic=False, name="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    p  = tf.paragraphs[0]; p.alignment = align
    r  = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold
    r.font.italic = italic; r.font.color.rgb = color; r.font.name = name
    return tb

def hbar(slide, y, color=ACCENT, thick=Inches(0.05)):
    rect(slide, 0, y, W, thick, color)

def tag(slide, label, x, y, color=CARD):
    rect(slide, x, y, Inches(1.9), Inches(0.38), color)
    txt(slide, label, x, y+Inches(0.05), Inches(1.9), Inches(0.3),
        size=9, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# ════════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE
# ════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
bg(s)

# Left indigo panel
rect(s, 0, 0, Inches(5.2), H, ACCENT)

# Brand
txt(s, "R E B O O T", Inches(0.4), Inches(0.45), Inches(4.5), Inches(0.5),
    size=11, bold=True, color=WHITE, name="Calibri")
rect(s, Inches(0.4), Inches(0.97), Inches(2.0), Inches(0.03), AMBER)

# Main title
txt(s, "OPTIMAL\nTRADE\nEXECUTION",
    Inches(0.4), Inches(1.15), Inches(4.5), Inches(3.0),
    size=48, bold=True, color=WHITE, name="Calibri")

# Tagline
txt(s, "Institutional-grade adaptive order execution\nwith real-time simulation & strategy benchmarking.",
    Inches(0.4), Inches(4.25), Inches(4.5), Inches(0.9),
    size=11, color=LIGHT, name="Calibri")

# Tags
for i, t in enumerate(["ADAPTIVE", "REAL-TIME", "SIMULATED"]):
    tag(s, t, Inches(0.4 + i*1.95), Inches(5.4), CARD)

txt(s, "Hackathon 2026  ·  Team Reboot",
    Inches(0.4), Inches(6.6), Inches(4.5), Inches(0.4),
    size=9, color=MUTED, name="Calibri")

# Right side — big stat
txt(s, "$500B+",
    Inches(5.6), Inches(1.6), Inches(7.0), Inches(1.8),
    size=80, bold=True, color=WHITE, align=PP_ALIGN.CENTER, name="Calibri")
txt(s, "lost annually to poor trade execution\nacross global institutional markets",
    Inches(5.6), Inches(3.4), Inches(7.0), Inches(0.9),
    size=14, color=MUTED, align=PP_ALIGN.CENTER, name="Calibri")

rect(s, Inches(6.5), Inches(4.5), Inches(5.0), Inches(0.03), CARD)

txt(s, "We built the system that fixes this.",
    Inches(5.6), Inches(4.7), Inches(7.0), Inches(0.6),
    size=15, bold=True, color=ACCENT, align=PP_ALIGN.CENTER, name="Calibri")


# ════════════════════════════════════════════════════════════════
# SLIDE 2 — THE PROBLEM (make them feel it)
# ════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
bg(s)

# Dark red tint overlay strip at top
rect(s, 0, 0, W, Inches(0.08), RED)

txt(s, "THE PROBLEM", Inches(0.5), Inches(0.2), Inches(8), Inches(0.55),
    size=30, bold=True, color=WHITE, name="Calibri")
txt(s, "Every large trade is a ticking clock — and the market knows you're there.",
    Inches(0.5), Inches(0.75), Inches(12), Inches(0.45),
    size=13, color=MUTED, name="Calibri")

rect(s, 0, Inches(1.18), W, Inches(0.03), CARD)

# ── Three Pain Stats (top row) ──────────────────────────────────
pain = [
    ("$0.10–$0.50",  "lost per share",     "on every large institutional order due to market impact alone.",      RED),
    ("30–50 bps",    "execution slippage", "is the average Implementation Shortfall on mid-cap equity orders.",   AMBER),
    ("3× worse",     "TWAP underperforms", "in volatile or illiquid markets vs an adaptive execution strategy.",  ACCENT),
]
for i, (stat, label, sub, color) in enumerate(pain):
    x = Inches(0.4 + i * 4.3)
    y = Inches(1.35)
    rect(s, x, y, Inches(4.0), Inches(2.3), CARD)
    rect(s, x, y, Inches(4.0), Inches(0.06), color)
    txt(s, stat, x+Inches(0.2), y+Inches(0.18),
        Inches(3.6), Inches(0.85),
        size=34, bold=True, color=color, name="Calibri")
    txt(s, label.upper(), x+Inches(0.2), y+Inches(1.02),
        Inches(3.6), Inches(0.38),
        size=10, bold=True, color=WHITE, name="Calibri")
    txt(s, sub, x+Inches(0.2), y+Inches(1.4),
        Inches(3.6), Inches(0.75),
        size=9, color=MUTED, name="Calibri")

# ── Central tension statement ────────────────────────────────────
rect(s, Inches(0.4), Inches(3.9), Inches(12.53), Inches(1.1), RGBColor(0x1a,0x1a,0x2e))
rect(s, Inches(0.4), Inches(3.9), Inches(0.07), Inches(1.1), RED)
txt(s,
    "A trader with 10,000 shares to sell has one window to act. Sell too fast → you crash the price "
    "against yourself. Sell too slow → the market moves away and you miss your window entirely. "
    "Static strategies like TWAP don't adapt — they just fail quietly.",
    Inches(0.65), Inches(3.98), Inches(12.0), Inches(0.95),
    size=11.5, color=LIGHT, name="Calibri")

# ── Three root causes ────────────────────────────────────────────
causes = [
    ("Market Impact",      "Each child order shifts prices against you — the larger the slice, the worse the price."),
    ("Regime Blindness",   "Volatility spikes and liquidity crises mid-execution break fixed-schedule strategies."),
    ("No Benchmarking",    "Without apples-to-apples comparison, traders can't know which strategy actually won."),
]
for i, (title, body) in enumerate(causes):
    x = Inches(0.4 + i * 4.3)
    y = Inches(5.25)
    rect(s, x, y, Inches(4.0), Inches(1.85), CARD)
    txt(s, f"↳  {title}", x+Inches(0.18), y+Inches(0.14),
        Inches(3.6), Inches(0.42),
        size=11, bold=True, color=AMBER, name="Calibri")
    txt(s, body, x+Inches(0.18), y+Inches(0.54),
        Inches(3.6), Inches(1.1),
        size=10, color=LIGHT, name="Calibri")


# ════════════════════════════════════════════════════════════════
# SLIDE 3 — OUR SOLUTION
# ════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
bg(s)
rect(s, 0, 0, W, Inches(0.08), TEAL)

txt(s, "OUR SOLUTION", Inches(0.5), Inches(0.2), Inches(8), Inches(0.55),
    size=30, bold=True, color=WHITE, name="Calibri")
txt(s, "A full-stack intelligent execution engine that adapts in real time.",
    Inches(0.5), Inches(0.75), Inches(12), Inches(0.4),
    size=13, color=MUTED, name="Calibri")
rect(s, 0, Inches(1.18), W, Inches(0.03), CARD)

# Central thesis
rect(s, Inches(0.4), Inches(1.3), Inches(12.53), Inches(0.95), ACCENT)
txt(s, "We simulate the market, replay the order, compare every strategy on the same data — "
       "and show exactly where and why one wins.",
    Inches(0.6), Inches(1.42), Inches(12.1), Inches(0.75),
    size=13, bold=False, color=WHITE, name="Calibri")

# Four pillars
pillars = [
    (ACCENT, "🔁  Market Simulator",
     "Geometric Brownian Motion price engine with configurable volatility, liquidity, spread, and volume. "
     "Supports 4 shock scenarios — volatility spike, liquidity crunch, combined crisis, and real Kaggle LOB data replay. "
     "Fully seed-driven for reproducible experiments."),
    (TEAL, "🧠  Adaptive Strategy Engine",
     "At every 1-second tick, the strategy reads the current market regime and adjusts its child-order size, "
     "aggressiveness, and participation rate. Detects volatility shocks and liquidity droughts in real time. "
     "Achieved 53% cost reduction vs TWAP baseline on Kaggle Optiver data."),
    (AMBER, "⚙️  Almgren-Chriss Cost Model",
     "Realistic execution cost decomposition: half-spread cost, square-root market impact (price permanent + temporary), "
     "and transaction fees — computed separately per BUY/SELL direction. Grounded in academic finance theory."),
    (RED, "📡  Live Streaming API",
     "FastAPI backend with WebSocket event bus streams every market tick, strategy decision, and fill result "
     "in real time to the React dashboard. REST endpoints for run creation, replay, and strategy benchmarking. "
     "PostgreSQL persistence for full audit trail."),
]

for i, (color, title, body) in enumerate(pillars):
    col, row = i % 2, i // 2
    x = Inches(0.4 + col * 6.5)
    y = Inches(2.45 + row * 2.35)
    rect(s, x, y, Inches(6.15), Inches(2.1), CARD)
    rect(s, x, y, Inches(0.07), Inches(2.1), color)
    txt(s, title, x+Inches(0.22), y+Inches(0.14),
        Inches(5.85), Inches(0.42),
        size=13, bold=True, color=WHITE, name="Calibri")
    txt(s, body, x+Inches(0.22), y+Inches(0.6),
        Inches(5.85), Inches(1.35),
        size=10, color=LIGHT, name="Calibri")


# ════════════════════════════════════════════════════════════════
# SLIDE 4 — ARCHITECTURE + FLOW
# ════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
bg(s)
rect(s, 0, 0, W, Inches(0.08), ACCENT)

txt(s, "ARCHITECTURE & EXECUTION FLOW", Inches(0.5), Inches(0.2), Inches(10), Inches(0.55),
    size=30, bold=True, color=WHITE, name="Calibri")
txt(s, "How the system is structured and how an order moves through it.",
    Inches(0.5), Inches(0.75), Inches(12), Inches(0.4),
    size=13, color=MUTED, name="Calibri")
rect(s, 0, Inches(1.18), W, Inches(0.03), CARD)

# ── LEFT: Architecture layers ─────────────────────────────────
layers = [
    ("PRESENTATION",  TEAL,   "React Dashboard · Live Charts · Config Form · Benchmark Panel"),
    ("APPLICATION",   ACCENT, "FastAPI REST  ·  WebSocket Event Bus  ·  Run Service  ·  Router"),
    ("CORE ENGINE",   AMBER,  "Market Simulator  ·  Fill Engine  ·  Strategy Engine  ·  Evaluator"),
    ("DATA",          MUTED,  "PostgreSQL / SQLite  ·  Kaggle LOB Replayer  ·  Scenario Configs"),
]
for i, (label, color, items) in enumerate(layers):
    y = Inches(1.38 + i * 1.42)
    rect(s, Inches(0.4), y, Inches(6.3), Inches(1.2), CARD)
    rect(s, Inches(0.4), y, Inches(0.07), Inches(1.2), color)
    txt(s, label, Inches(0.62), y+Inches(0.1),
        Inches(3.0), Inches(0.35),
        size=9, bold=True, color=color, name="Calibri")
    txt(s, items, Inches(0.62), y+Inches(0.45),
        Inches(6.0), Inches(0.65),
        size=10.5, color=LIGHT, name="Calibri")
    if i < 3:
        txt(s, "▼", Inches(3.2), y+Inches(1.2),
            Inches(0.5), Inches(0.3),
            size=12, color=MUTED, align=PP_ALIGN.CENTER, name="Calibri")

# ── RIGHT: Execution flow steps ───────────────────────────────
flow_steps = [
    (ACCENT, "01  Configure",   "Symbol · Side · Qty · Strategy · Scenario · Seed"),
    (TEAL,   "02  Simulate",    "GBM price path + shock injection at t=30s"),
    (AMBER,  "03  Decide",      "Strategy slices order into optimal child-order per tick"),
    (TEAL,   "04  Fill",        "Volume cap → spread + impact + fee cost computed"),
    (ACCENT, "05  Stream",      "WebSocket broadcasts market state + fill live"),
    (RED,    "06  Evaluate",    "VWAP · IS · Cost (bps) · Risk Score · Stability"),
    (MUTED,  "07  Benchmark",   "All 3 strategies replayed on identical market path"),
]
for i, (color, title, sub) in enumerate(flow_steps):
    y = Inches(1.38 + i * 0.87)
    rect(s, Inches(7.1), y, Inches(5.9), Inches(0.75), CARD)
    rect(s, Inches(7.1), y, Inches(0.07), Inches(0.75), color)
    txt(s, title, Inches(7.32), y+Inches(0.06),
        Inches(5.5), Inches(0.3),
        size=11, bold=True, color=WHITE, name="Calibri")
    txt(s, sub, Inches(7.32), y+Inches(0.36),
        Inches(5.5), Inches(0.3),
        size=9.5, color=MUTED, name="Calibri")
    if i < 6:
        txt(s, "↓", Inches(9.8), y+Inches(0.75),
            Inches(0.3), Inches(0.22),
            size=9, color=CARD, align=PP_ALIGN.CENTER, name="Calibri")

# divider between left and right
rect(s, Inches(6.9), Inches(1.35), Inches(0.03), Inches(5.8), CARD)


# ════════════════════════════════════════════════════════════════
# SLIDE 5 — TECH STACK + RESULTS
# ════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
bg(s)
rect(s, 0, 0, W, Inches(0.08), AMBER)

txt(s, "TECH STACK & RESULTS", Inches(0.5), Inches(0.2), Inches(10), Inches(0.55),
    size=30, bold=True, color=WHITE, name="Calibri")
txt(s, "What we built with — and proof that it works.",
    Inches(0.5), Inches(0.75), Inches(12), Inches(0.4),
    size=13, color=MUTED, name="Calibri")
rect(s, 0, Inches(1.18), W, Inches(0.03), CARD)

# ── LEFT: Tech Stack ─────────────────────────────────────────
txt(s, "TECH STACK", Inches(0.5), Inches(1.35), Inches(5), Inches(0.38),
    size=9, bold=True, color=MUTED, name="Calibri")

stack = [
    (ACCENT, "FastAPI",         "REST API · WebSocket streaming · async routes"),
    (TEAL,   "Pydantic v2",     "Schema validation · request/response models"),
    (AMBER,  "SQLAlchemy 2.x",  "ORM · PostgreSQL + SQLite support · async sessions"),
    (MUTED,  "NumPy / Pandas",  "GBM simulation · Kaggle LOB data processing"),
    (RED,    "WebSockets",      "Live tick streaming to React frontend"),
    (ACCENT, "React",           "Dashboard · charts · config form (Frontend layer)"),
]
for i, (color, name_, desc) in enumerate(stack):
    y = Inches(1.75 + i * 0.87)
    rect(s, Inches(0.4), y, Inches(6.0), Inches(0.73), CARD)
    rect(s, Inches(0.4), y, Inches(0.07), Inches(0.73), color)
    txt(s, name_, Inches(0.62), y+Inches(0.07),
        Inches(5.5), Inches(0.3),
        size=12, bold=True, color=WHITE, name="Calibri")
    txt(s, desc, Inches(0.62), y+Inches(0.38),
        Inches(5.5), Inches(0.3),
        size=9, color=MUTED, name="Calibri")

# ── RIGHT: Results ───────────────────────────────────────────
rect(s, Inches(6.9), Inches(1.35), Inches(0.03), Inches(5.8), CARD)

txt(s, "BENCHMARK RESULTS  (Kaggle Optiver LOB · Vol Shock · 10,000 shares)",
    Inches(7.1), Inches(1.35), Inches(5.9), Inches(0.4),
    size=9, bold=True, color=MUTED, name="Calibri")

# Table header
rect(s, Inches(7.1), Inches(1.78), Inches(5.9), Inches(0.42), ACCENT)
for j, h in enumerate(["Strategy", "Cost ($)", "Cost (bps)", "Stability"]):
    txt(s, h, Inches(7.15 + j * 1.46), Inches(1.85),
        Inches(1.4), Inches(0.28),
        size=9, bold=True, color=WHITE, name="Calibri")

rows = [
    ("TWAP",          "$3,182",  "31.8 bps",  "0.99", CARD, WHITE),
    ("Volume-Aware",  "$1,547",  "15.5 bps",  "0.97", CARD, WHITE),
    ("Adaptive ★",   "$1,484",  "14.8 bps",  "0.96", RGBColor(0x1e,0x1b,0x4b), TEAL),
]
for i, (name_, cost, bps, stab, bg_c, fc) in enumerate(rows):
    ry = Inches(2.2 + i * 0.6)
    rect(s, Inches(7.1), ry, Inches(5.9), Inches(0.55), bg_c)
    for j, val in enumerate([name_, cost, bps, stab]):
        txt(s, val, Inches(7.15 + j * 1.46), ry+Inches(0.12),
            Inches(1.4), Inches(0.32),
            size=11, bold=(i == 2), color=fc, name="Calibri")

# Big callout
rect(s, Inches(7.1), Inches(3.65), Inches(5.9), Inches(1.5), ACCENT)
txt(s, "53%", Inches(7.1), Inches(3.72),
    Inches(5.9), Inches(0.95),
    size=62, bold=True, color=WHITE, align=PP_ALIGN.CENTER, name="Calibri")
txt(s, "cost reduction vs TWAP  ·  same market · same data · same seed",
    Inches(7.1), Inches(4.6), Inches(5.9), Inches(0.42),
    size=9.5, color=LIGHT, align=PP_ALIGN.CENTER, name="Calibri")

# Test coverage row
txt(s, "TEST COVERAGE", Inches(7.1), Inches(5.35), Inches(5.9), Inches(0.35),
    size=9, bold=True, color=MUTED, name="Calibri")

for i, (num, label, color) in enumerate([
    ("131", "Total Tests",   ACCENT),
    ("81",  "Backend",       TEAL),
    ("46",  "Engine",        AMBER),
    ("4",   "Kaggle",        RED),
]):
    bx = Inches(7.1 + i * 1.47)
    by = Inches(5.7)
    rect(s, bx, by, Inches(1.35), Inches(1.3), CARD)
    rect(s, bx, by, Inches(1.35), Inches(0.05), color)
    txt(s, num, bx+Inches(0.1), by+Inches(0.12),
        Inches(1.15), Inches(0.7),
        size=30, bold=True, color=color, align=PP_ALIGN.CENTER, name="Calibri")
    txt(s, label, bx+Inches(0.1), by+Inches(0.85),
        Inches(1.15), Inches(0.35),
        size=8.5, color=MUTED, align=PP_ALIGN.CENTER, name="Calibri")


# ════════════════════════════════════════════════════════════════
# Save
prs.save("Reboot_Optimal_Trade_Execution.pptx")
print(f"✓  Saved — {len(prs.slides)} slides")
